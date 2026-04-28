from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, distinct
from typing import List, Optional
import io
import json
from datetime import date, datetime
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from docx import Document as DocxDocument
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.database import get_db
from app.models import (
    ReportTemplate, User, Project, Beneficiary, Transaction, Indicator,
    Measurement, Distribution, DistributionItem, CashTransfer,
    DataCollectionForm, FormSubmission, ReportType, IPTTEntry,
    Complaint, ComplaintStatus, FieldVisit, ComplianceAssessment,
    Recommendation, RecommendationStatus, LessonLearned,
)
from app.schemas import (
    ReportTemplateCreate, ReportTemplateOut, ReportGenerateRequest,
)
from app.auth import get_current_user

router = APIRouter(prefix="/api/reports", tags=["التقارير"])


@router.get("/templates", response_model=List[ReportTemplateOut])
def list_templates(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return db.query(ReportTemplate).order_by(ReportTemplate.created_at.desc()).all()


@router.post("/templates", response_model=ReportTemplateOut)
def create_template(
    data: ReportTemplateCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    tmpl = ReportTemplate(
        name=data.name,
        description=data.description,
        report_type=data.report_type,
        template_config=data.template_config,
        created_by=current_user.id,
    )
    db.add(tmpl)
    db.commit()
    db.refresh(tmpl)
    return tmpl


@router.delete("/templates/{template_id}")
def delete_template(
    template_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    tmpl = db.query(ReportTemplate).filter(ReportTemplate.id == template_id).first()
    if not tmpl:
        raise HTTPException(status_code=404, detail="القالب غير موجود")
    db.delete(tmpl)
    db.commit()
    return {"message": "تم حذف القالب بنجاح"}


def _style_excel_header(ws, headers, row=1):
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    return row + 1


def _generate_project_progress_excel(db, project_id, start_date, end_date):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "ملخص المشاريع"
    ws1.sheet_view.rightToLeft = True

    query = db.query(Project)
    if project_id:
        query = query.filter(Project.id == project_id)
    projects = query.all()

    headers = ["رمز المشروع", "اسم المشروع", "القطاع", "الحالة", "الميزانية", "المصروف",
               "نسبة الإنفاق", "المستفيدون المستهدفون", "المستفيدون الفعليون", "المحافظة", "المانح"]
    row = _style_excel_header(ws1, headers)

    for p in projects:
        spent_pct = (p.spent / p.budget * 100) if p.budget > 0 else 0
        ws1.append([
            p.code, p.name, p.sector, p.status.value if p.status else "",
            p.budget, p.spent, f"{spent_pct:.1f}%",
            p.target_beneficiaries, p.actual_beneficiaries,
            p.governorate, p.donor,
        ])

    for col in ws1.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws1.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    ws2 = wb.create_sheet("المؤشرات")
    ws2.sheet_view.rightToLeft = True
    ind_headers = ["رمز المؤشر", "المؤشر", "النوع", "المستهدف", "الفعلي", "نسبة الإنجاز", "التكرار", "المشروع"]
    row = _style_excel_header(ws2, ind_headers)

    ind_query = db.query(Indicator)
    if project_id:
        ind_query = ind_query.filter(Indicator.project_id == project_id)
    indicators = ind_query.all()
    for ind in indicators:
        pct = (ind.actual_value / ind.target_value * 100) if ind.target_value > 0 else 0
        project = db.query(Project).filter(Project.id == ind.project_id).first()
        ws2.append([
            ind.code, ind.name, ind.type.value if ind.type else "",
            ind.target_value, ind.actual_value, f"{pct:.1f}%",
            ind.frequency, project.name if project else "",
        ])

    for col in ws2.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws2.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    return wb


def _generate_beneficiary_list_excel(db, project_id, governorate):
    wb = Workbook()
    ws = wb.active
    ws.title = "قائمة المستفيدين"
    ws.sheet_view.rightToLeft = True

    headers = ["#", "الرقم الوطني", "الاسم الأول", "الاسم الأخير", "الجنس",
               "الهاتف", "المحافظة", "المديرية", "حجم الأسرة", "درجة الضعف", "الحالة"]
    row = _style_excel_header(ws, headers)

    query = db.query(Beneficiary)
    if governorate:
        query = query.filter(Beneficiary.governorate == governorate)
    beneficiaries = query.order_by(Beneficiary.created_at.desc()).all()

    for i, b in enumerate(beneficiaries, 1):
        ws.append([
            i, b.national_id, b.first_name, b.last_name,
            b.gender.value if b.gender else "",
            b.phone, b.governorate, b.district,
            b.household_size, b.vulnerability_score,
            b.status.value if b.status else "",
        ])

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    return wb


def _generate_financial_summary_excel(db, project_id, start_date, end_date):
    wb = Workbook()
    ws = wb.active
    ws.title = "الملخص المالي"
    ws.sheet_view.rightToLeft = True

    headers = ["المرجع", "النوع", "المبلغ", "العملة", "الوصف", "الفئة", "التاريخ"]
    row = _style_excel_header(ws, headers)

    query = db.query(Transaction)
    if project_id:
        query = query.filter(Transaction.project_id == project_id)
    if start_date:
        query = query.filter(Transaction.transaction_date >= start_date)
    if end_date:
        query = query.filter(Transaction.transaction_date <= end_date)
    transactions = query.order_by(Transaction.transaction_date.desc()).all()

    for t in transactions:
        ws.append([
            t.reference, t.type.value if t.type else "",
            t.amount, t.currency.value if t.currency else "",
            t.description, t.category,
            str(t.transaction_date) if t.transaction_date else "",
        ])

    total_income = sum(t.amount for t in transactions if t.type and t.type.value == "income")
    total_expense = sum(t.amount for t in transactions if t.type and t.type.value == "expense")
    ws.append([])
    ws.append(["", "إجمالي الإيرادات", total_income])
    ws.append(["", "إجمالي المصروفات", total_expense])
    ws.append(["", "الصافي", total_income - total_expense])

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    return wb


def _generate_survey_analysis_excel(db, form_id):
    wb = Workbook()
    ws = wb.active
    ws.title = "تحليل الاستجابات"
    ws.sheet_view.rightToLeft = True

    form = db.query(DataCollectionForm).filter(DataCollectionForm.id == form_id).first()
    if not form:
        headers = ["لا توجد بيانات"]
        _style_excel_header(ws, headers)
        return wb

    field_labels = [f.label for f in form.fields]
    headers = ["#", "المحافظة", "المديرية", "التاريخ", "الحالة"] + field_labels
    _style_excel_header(ws, headers)

    submissions = db.query(FormSubmission).filter(FormSubmission.form_id == form_id).all()
    for i, sub in enumerate(submissions, 1):
        try:
            data = json.loads(sub.data)
        except json.JSONDecodeError:
            data = {}
        row_data = [
            i, sub.governorate or "", sub.district or "",
            str(sub.submitted_at) if sub.submitted_at else "",
            sub.status.value if sub.status else "",
        ]
        for field in form.fields:
            row_data.append(str(data.get(field.field_name, "")))
        ws.append(row_data)

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    return wb


def _generate_project_progress_docx(db, project_id, title):
    doc = DocxDocument()

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    doc.add_heading(title or "تقرير تقدم المشاريع", level=0)
    doc.add_paragraph(f"تاريخ التقرير: {date.today().isoformat()}")
    doc.add_paragraph("")

    query = db.query(Project)
    if project_id:
        query = query.filter(Project.id == project_id)
    projects = query.all()

    for p in projects:
        doc.add_heading(f"{p.code} - {p.name}", level=1)

        table = doc.add_table(rows=8, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        fields = [
            ("القطاع", p.sector or "-"),
            ("الحالة", p.status.value if p.status else "-"),
            ("الميزانية", f"{p.budget:,.2f} {p.currency.value}" if p.currency else str(p.budget)),
            ("المصروف", f"{p.spent:,.2f}"),
            ("المستفيدون المستهدفون", str(p.target_beneficiaries)),
            ("المستفيدون الفعليون", str(p.actual_beneficiaries)),
            ("المحافظة", p.governorate or "-"),
            ("المانح", p.donor or "-"),
        ]
        for i, (label, value) in enumerate(fields):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = str(value)
            for cell in table.rows[i].cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_paragraph("")

        indicators = db.query(Indicator).filter(Indicator.project_id == p.id).all()
        if indicators:
            doc.add_heading("المؤشرات", level=2)
            ind_table = doc.add_table(rows=1, cols=5)
            ind_table.style = "Table Grid"
            hdr = ind_table.rows[0].cells
            hdr[0].text = "المؤشر"
            hdr[1].text = "النوع"
            hdr[2].text = "المستهدف"
            hdr[3].text = "الفعلي"
            hdr[4].text = "نسبة الإنجاز"
            for cell in hdr:
                for paragraph in cell.paragraphs:
                    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                    run.bold = True

            for ind in indicators:
                pct = (ind.actual_value / ind.target_value * 100) if ind.target_value > 0 else 0
                row = ind_table.add_row().cells
                row[0].text = ind.name
                row[1].text = ind.type.value if ind.type else ""
                row[2].text = str(ind.target_value)
                row[3].text = str(ind.actual_value)
                row[4].text = f"{pct:.1f}%"

        doc.add_page_break()

    return doc


@router.post("/generate")
def generate_report(
    data: ReportGenerateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    report_title = data.title or "تقرير"

    if data.format == "excel":
        if data.report_type == ReportType.PROJECT_PROGRESS:
            wb = _generate_project_progress_excel(db, data.project_id, data.start_date, data.end_date)
        elif data.report_type == ReportType.BENEFICIARY_LIST:
            wb = _generate_beneficiary_list_excel(db, data.project_id, data.governorate)
        elif data.report_type == ReportType.FINANCIAL_SUMMARY:
            wb = _generate_financial_summary_excel(db, data.project_id, data.start_date, data.end_date)
        elif data.report_type == ReportType.SURVEY_ANALYSIS:
            wb = _generate_survey_analysis_excel(db, data.form_id)
        elif data.report_type == ReportType.INDICATOR_TRACKING:
            wb = _generate_project_progress_excel(db, data.project_id, data.start_date, data.end_date)
        else:
            wb = _generate_project_progress_excel(db, data.project_id, data.start_date, data.end_date)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"{report_title}_{date.today().isoformat()}.xlsx"
        encoded_filename = quote(filename)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
        )

    elif data.format == "word":
        if data.report_type == ReportType.PROJECT_PROGRESS:
            doc = _generate_project_progress_docx(db, data.project_id, report_title)
        else:
            doc = _generate_project_progress_docx(db, data.project_id, report_title)

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        filename = f"{report_title}_{date.today().isoformat()}.docx"
        encoded_filename = quote(filename)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
        )

    else:
        raise HTTPException(status_code=400, detail="صيغة غير مدعومة. استخدم excel أو word")


@router.get("/types")
def list_report_types(current_user: User = Depends(get_current_user)):
    return [
        {"value": "project_progress", "label": "تقرير تقدم المشاريع"},
        {"value": "beneficiary_list", "label": "قائمة المستفيدين"},
        {"value": "financial_summary", "label": "الملخص المالي"},
        {"value": "indicator_tracking", "label": "تتبع المؤشرات"},
        {"value": "distribution_report", "label": "تقرير التوزيعات"},
        {"value": "survey_analysis", "label": "تحليل الاستبيانات"},
    ]


# -- Cluster & Donor Report Templates --

CLUSTER_TEMPLATES = {
    "WASH": {"name": "تقرير قطاع المياه والصرف الصحي", "sections": ["الأنشطة المنفذة", "المستفيدون", "المؤشرات", "التحديات", "التوصيات"]},
    "FSL": {"name": "تقرير قطاع الأمن الغذائي", "sections": ["التوزيعات الغذائية", "برامج سبل العيش", "المستفيدون", "الأمن الغذائي", "التحديات"]},
    "Protection": {"name": "تقرير قطاع الحماية", "sections": ["حالات الحماية", "الإحالات", "التوعية", "الأنشطة النفسية", "التحديات"]},
    "Health": {"name": "تقرير القطاع الصحي", "sections": ["الخدمات الصحية", "التطعيمات", "التغذية", "صحة الأم والطفل", "الأوبئة"]},
    "Education": {"name": "تقرير قطاع التعليم", "sections": ["الالتحاق", "المدارس المدعومة", "المعلمون", "المستلزمات", "التحديات"]},
    "Shelter": {"name": "تقرير قطاع المأوى", "sections": ["المآوي المقدمة", "الترميمات", "المواد", "المستفيدون", "التحديات"]},
}

DONOR_TEMPLATES = {
    "USAID": {"name": "تقرير USAID", "sections": ["ملخص تنفيذي", "تقدم المؤشرات", "النتائج الرئيسية", "الإنفاق المالي", "الدروس المستفادة", "الخطة القادمة"]},
    "ECHO": {"name": "تقرير ECHO", "sections": ["نظرة عامة", "الأنشطة والنتائج", "المستفيدون", "الميزانية", "التنسيق", "الاستدامة"]},
    "OCHA": {"name": "تقرير OCHA", "sections": ["الملخص", "الاحتياجات", "الاستجابة", "الفجوات", "التمويل", "التنسيق"]},
    "UNICEF": {"name": "تقرير UNICEF", "sections": ["ملخص تنفيذي", "نتائج البرنامج", "الأطفال المستفيدون", "المالية", "المخاطر", "الخطوات القادمة"]},
    "WFP": {"name": "تقرير WFP", "sections": ["التوزيعات", "المستفيدون", "سلسلة الإمداد", "المراقبة", "التمويل"]},
}


@router.get("/templates/cluster")
def get_cluster_templates(current_user: User = Depends(get_current_user)):
    return CLUSTER_TEMPLATES


@router.get("/templates/donor")
def get_donor_templates(current_user: User = Depends(get_current_user)):
    return DONOR_TEMPLATES


@router.get("/generate-cluster/{sector}")
def generate_cluster_report(
    sector: str,
    project_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    template = CLUSTER_TEMPLATES.get(sector)
    if not template:
        raise HTTPException(status_code=404, detail="القالب غير موجود")

    query = db.query(Project)
    if project_id:
        query = query.filter(Project.id == project_id)
    if sector != "all":
        query = query.filter(Project.sector == sector)
    projects = query.all()

    total_beneficiaries = 0
    total_budget = 0
    project_data = []
    for p in projects:
        bcount = db.query(func.count(distinct(DistributionItem.beneficiary_id))).join(
            Distribution, DistributionItem.distribution_id == Distribution.id
        ).filter(Distribution.project_id == p.id).scalar() or 0
        total_beneficiaries += bcount
        total_budget += p.budget or 0
        project_data.append({"name": p.name, "beneficiaries": bcount, "budget": p.budget or 0, "governorate": p.governorate})

    return {
        "template": template,
        "sector": sector,
        "generated_at": datetime.utcnow().isoformat(),
        "summary": {
            "total_projects": len(projects),
            "total_beneficiaries": total_beneficiaries,
            "total_budget": total_budget,
        },
        "projects": project_data,
    }


@router.get("/generate-monthly-meal/{project_id}")
def generate_monthly_meal_report(
    project_id: int,
    year: Optional[int] = None,
    month: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="المشروع غير موجود")

    indicators = db.query(Indicator).filter(Indicator.project_id == project_id).all()
    iptt_q = db.query(IPTTEntry).filter(IPTTEntry.project_id == project_id)
    if year:
        iptt_q = iptt_q.filter(IPTTEntry.year == year)
    if month:
        iptt_q = iptt_q.filter(IPTTEntry.month == month)
    iptt_entries = iptt_q.all()

    complaints = db.query(Complaint).filter(Complaint.project_id == project_id).all()
    visits = db.query(FieldVisit).filter(FieldVisit.project_id == project_id).all()
    lessons = db.query(LessonLearned).filter(LessonLearned.project_id == project_id).all()
    recommendations = db.query(Recommendation).filter(Recommendation.project_id == project_id).all()

    green = sum(1 for e in iptt_entries if e.status_color == "green")
    yellow = sum(1 for e in iptt_entries if e.status_color == "yellow")
    red = sum(1 for e in iptt_entries if e.status_color == "red")

    return {
        "report_type": "monthly_meal",
        "project": {"id": project.id, "name": project.name, "sector": project.sector},
        "generated_at": datetime.utcnow().isoformat(),
        "monitoring": {
            "total_indicators": len(indicators),
            "iptt_entries": len(iptt_entries),
            "performance": {"green": green, "yellow": yellow, "red": red},
        },
        "evaluation": {
            "field_visits": len(visits),
        },
        "accountability": {
            "total_complaints": len(complaints),
            "resolved": sum(1 for c in complaints if c.status in [ComplaintStatus.RESOLVED, ComplaintStatus.CLOSED]),
            "pending": sum(1 for c in complaints if c.status in [ComplaintStatus.RECEIVED, ComplaintStatus.IN_PROGRESS]),
        },
        "learning": {
            "lessons_learned": len(lessons),
            "total_recommendations": len(recommendations),
            "completed_recommendations": sum(1 for r in recommendations if r.status == RecommendationStatus.COMPLETED),
        },
    }


@router.get("/generate-iptt/{project_id}")
def generate_iptt_report(
    project_id: int,
    year: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="المشروع غير موجود")

    indicators = db.query(Indicator).filter(Indicator.project_id == project_id).all()
    result = []
    for ind in indicators:
        q = db.query(IPTTEntry).filter(IPTTEntry.indicator_id == ind.id)
        if year:
            q = q.filter(IPTTEntry.year == year)
        entries = q.order_by(IPTTEntry.year, IPTTEntry.month).all()
        latest = entries[-1] if entries else None
        result.append({
            "code": ind.code, "name": ind.name, "type": ind.type.value if ind.type else None,
            "unit": ind.unit, "baseline": ind.baseline, "annual_target": ind.target_value,
            "cumulative_actual": latest.cumulative_actual if latest else 0,
            "achievement_rate": latest.achievement_rate if latest else 0,
            "status_color": latest.status_color if latest else "green",
            "monthly": [{"month": e.month, "target": e.target_value, "actual": e.actual_value} for e in entries],
        })

    return {
        "report_type": "iptt",
        "project": {"id": project.id, "name": project.name},
        "generated_at": datetime.utcnow().isoformat(),
        "indicators": result,
    }


@router.get("/generate-cfm")
def generate_cfm_report(
    project_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Complaint)
    if project_id:
        query = query.filter(Complaint.project_id == project_id)
    complaints = query.all()

    by_status = {}
    by_channel = {}
    by_category = {}
    for c in complaints:
        s = c.status.value if c.status else "unknown"
        by_status[s] = by_status.get(s, 0) + 1
        ch = c.channel.value if c.channel else "unknown"
        by_channel[ch] = by_channel.get(ch, 0) + 1
        cat = c.category.value if c.category else "unknown"
        by_category[cat] = by_category.get(cat, 0) + 1

    resolved = [c for c in complaints if c.resolution_date and c.created_at]
    avg_days = round(
        sum((c.resolution_date - c.created_at).days for c in resolved) / len(resolved), 1
    ) if resolved else 0

    return {
        "report_type": "cfm",
        "generated_at": datetime.utcnow().isoformat(),
        "total_complaints": len(complaints),
        "by_status": by_status,
        "by_channel": by_channel,
        "by_category": by_category,
        "avg_resolution_days": avg_days,
        "sensitive_cases": sum(1 for c in complaints if c.is_sensitive),
    }


@router.get("/generate-compliance/{project_id}")
def generate_compliance_report(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    assessments = db.query(ComplianceAssessment).filter(
        ComplianceAssessment.project_id == project_id
    ).all()

    by_area = {}
    for a in assessments:
        area = a.area.value if a.area else "unknown"
        if area not in by_area:
            by_area[area] = {"total": 0, "compliant": 0, "partial": 0, "non_compliant": 0, "scores": []}
        by_area[area]["total"] += 1
        by_area[area]["scores"].append(a.score)
        st = a.status.value if a.status else ""
        if st == "compliant":
            by_area[area]["compliant"] += 1
        elif st == "partially_compliant":
            by_area[area]["partial"] += 1
        elif st == "non_compliant":
            by_area[area]["non_compliant"] += 1

    for area in by_area:
        scores = by_area[area]["scores"]
        by_area[area]["avg_score"] = round(sum(scores) / len(scores), 1) if scores else 0
        del by_area[area]["scores"]

    return {
        "report_type": "compliance",
        "project_id": project_id,
        "generated_at": datetime.utcnow().isoformat(),
        "total_assessments": len(assessments),
        "by_area": by_area,
    }
